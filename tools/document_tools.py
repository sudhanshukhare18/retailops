import os

from docx import Document


# ============================================================
# PATH
# ============================================================

BASE_DIR = os.path.dirname(
    os.path.dirname(
        os.path.abspath(__file__)
    )
)

TEMPLATE_PATH = os.path.join(
    BASE_DIR,
    "resources",
    "documents",
    "bill_template.docx"
)


# ============================================================
# REPLACE PLACEHOLDERS
# ============================================================

def replace_placeholders(
    document,
    replacements
):

    # --------------------------------------------------------
    # Paragraphs
    # --------------------------------------------------------

    for paragraph in document.paragraphs:

        for placeholder, value in replacements.items():

            if placeholder in paragraph.text:

                for run in paragraph.runs:

                    run.text = run.text.replace(
                        placeholder,
                        str(value)
                    )


    # --------------------------------------------------------
    # Tables
    # --------------------------------------------------------

    for table in document.tables:

        for row in table.rows:

            for cell in row.cells:

                for paragraph in cell.paragraphs:

                    for placeholder, value in replacements.items():

                        if placeholder in paragraph.text:

                            for run in paragraph.runs:

                                run.text = run.text.replace(
                                    placeholder,
                                    str(value)
                                )


# ============================================================
# GENERATE BILL DOCUMENT
# ============================================================

def generate_bill_document(
    bill_data,
    output_path
):

    if not os.path.exists(TEMPLATE_PATH):

        return {
            "success": False,
            "message": (
                "Bill template not found: "
                f"{TEMPLATE_PATH}"
            )
        }


    # --------------------------------------------------------
    # SECURITY:
    # Only customer-safe fields are accepted.
    # --------------------------------------------------------

    allowed_fields = {
        "customer_name",
        "order_id",
        "order_date",
        "items_text",
        "subtotal",
        "discount",
        "final_amount",
        "delivery_status"
    }


    safe_bill_data = {

        key: value

        for key, value in bill_data.items()

        if key in allowed_fields
    }


    # --------------------------------------------------------
    # Load DOCX template
    # --------------------------------------------------------

    document = Document(
        TEMPLATE_PATH
    )


    # --------------------------------------------------------
    # Placeholder values
    # --------------------------------------------------------

    replacements = {

        "{{customer_name}}":
            safe_bill_data.get(
                "customer_name",
                ""
            ),

        "{{order_id}}":
            safe_bill_data.get(
                "order_id",
                ""
            ),

        "{{order_date}}":
            safe_bill_data.get(
                "order_date",
                ""
            ),

        "{{items}}":
            safe_bill_data.get(
                "items_text",
                ""
            ),

        "{{subtotal}}":
            safe_bill_data.get(
                "subtotal",
                0
            ),

        "{{discount}}":
            safe_bill_data.get(
                "discount",
                0
            ),

        "{{final_amount}}":
            safe_bill_data.get(
                "final_amount",
                0
            ),

        "{{delivery_status}}":
            safe_bill_data.get(
                "delivery_status",
                ""
            )
    }


    # --------------------------------------------------------
    # Replace values
    # --------------------------------------------------------

    replace_placeholders(
        document,
        replacements
    )


    # --------------------------------------------------------
    # Save generated document
    # --------------------------------------------------------

    document.save(
        output_path
    )


    return {
        "success": True,
        "file": output_path,
        "message": "Customer bill document generated successfully."
    }